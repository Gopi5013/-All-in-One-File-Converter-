import re
import sys
from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from PIL import Image

from app import create_app


@pytest.fixture
def app(tmp_path):
    app = create_app()
    app.config.update(
        TESTING=True,
        SECRET_KEY="test-secret",
        OUTPUT_FOLDER=tmp_path / "outputs",
    )
    return app


@pytest.fixture
def client(app):
    return app.test_client()


def pdf_file(page_sizes=((100, 200),)):
    writer = PdfWriter()
    for width, height in page_sizes:
        writer.add_blank_page(width=width, height=height)
    file = BytesIO()
    writer.write(file)
    writer.close()
    file.seek(0)
    return file


def post_merge(client, files):
    return client.post(
        "/tools/merge-pdf",
        data={"pdfs": files},
        content_type="multipart/form-data",
    )


def post_pdf_to_word(client, uploaded_file):
    return client.post(
        "/tools/pdf-to-word",
        data={"pdf": uploaded_file},
        content_type="multipart/form-data",
    )


def download_url(response):
    match = re.search(
        r'href="(/tools/merge-pdf/[^"]+/download)"',
        response.get_data(as_text=True),
    )
    assert match is not None
    return match.group(1)


def pdf_to_word_download_url(response):
    match = re.search(
        r'href="(/tools/pdf-to-word/[^"]+/download)"',
        response.get_data(as_text=True),
    )
    assert match is not None
    return match.group(1)


def operation_download_url(response, operation):
    match = re.search(
        rf'href="(/tools/pdf-download/{re.escape(operation)}/[^"]+)"',
        response.get_data(as_text=True),
    )
    assert match is not None
    return match.group(1)


def image_file(size=(100, 60), image_format="PNG"):
    file = BytesIO()
    Image.new("RGB", size, color="navy").save(file, format=image_format)
    file.seek(0)
    return file


def test_merge_pdfs_preserves_order_and_download_is_session_scoped(app, client):
    response = post_merge(
        client,
        [
            (pdf_file(((100, 200),)), "first.pdf"),
            (pdf_file(((300, 400), (500, 600))), "second.pdf"),
        ],
    )

    assert response.status_code == 200
    assert b"merged into one document with 3 pages" in response.data

    url = download_url(response)
    download = client.get(url)
    assert download.status_code == 200
    assert download.mimetype == "application/pdf"
    assert download.headers["Content-Disposition"].endswith('filename=merged.pdf')

    merged = PdfReader(BytesIO(download.data), strict=True)
    assert len(merged.pages) == 3
    assert [float(page.mediabox.width) for page in merged.pages] == [
        100.0,
        300.0,
        500.0,
    ]

    other_client = app.test_client()
    assert other_client.get(url).status_code == 404


def test_homepage_marks_pdf_merger_as_available(client):
    response = client.get("/")

    assert response.status_code == 200
    assert b'href="/tools/pdf-to-word"' in response.data
    assert b'href="/tools/merge-pdf"' in response.data
    assert b"13</strong> available of 13 tools" in response.data


def test_merge_requires_at_least_two_pdfs(client):
    response = post_merge(client, [(pdf_file(), "only.pdf")])

    assert response.status_code == 400
    assert b"Choose at least two PDF files to merge." in response.data


def test_non_pdf_extension_is_rejected(client):
    response = post_merge(
        client,
        [
            (pdf_file(), "first.pdf"),
            (BytesIO(b"not a pdf"), "notes.txt"),
        ],
    )

    assert response.status_code == 400
    assert b"Choose PDF files only." in response.data


def test_corrupt_pdf_is_rejected(client):
    response = post_merge(
        client,
        [
            (pdf_file(), "first.pdf"),
            (BytesIO(b"not a pdf"), "corrupt.pdf"),
        ],
    )

    assert response.status_code == 400
    assert b"not a valid, readable PDF" in response.data


def test_empty_pdf_is_rejected(client):
    response = post_merge(
        client,
        [
            (pdf_file(), "first.pdf"),
            (BytesIO(), "empty.pdf"),
        ],
    )

    assert response.status_code == 400
    assert b"not a valid, readable PDF" in response.data


def test_large_pdf_request_returns_clear_error(app):
    app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024
    client = app.test_client()

    response = post_merge(
        client,
        [
            (BytesIO(b"x" * (1024 * 1024)), "first.pdf"),
            (pdf_file(), "second.pdf"),
        ],
    )

    assert response.status_code == 413
    assert b"maximum total upload size is 1 MB" in response.data


def test_pdf_to_word_conversion_and_download_are_session_scoped(app, client):
    response = post_pdf_to_word(client, (pdf_file(), "source.pdf"))

    assert response.status_code == 200
    assert b"Converted 1 PDF page to a Word document." in response.data

    url = pdf_to_word_download_url(response)
    download = client.get(url)
    assert download.status_code == 200
    assert download.mimetype == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert download.headers["Content-Disposition"].endswith(
        "filename=converted.docx"
    )
    document = Document(BytesIO(download.data))
    assert len(document.sections) == 1

    other_client = app.test_client()
    assert other_client.get(url).status_code == 404


def test_pdf_to_word_page_is_available_from_homepage(client):
    response = client.get("/tools/pdf-to-word")

    assert response.status_code == 200
    assert b"Scanned PDFs need OCR" in response.data
    assert b'id="pdf-to-word-form"' in response.data


def test_pdf_to_word_rejects_missing_or_unsupported_upload(client):
    missing_response = post_pdf_to_word(client, (BytesIO(), ""))
    unsupported_response = post_pdf_to_word(
        client,
        (BytesIO(b"not a PDF"), "notes.txt"),
    )

    assert missing_response.status_code == 400
    assert b"Choose a PDF file before converting." in missing_response.data
    assert unsupported_response.status_code == 400
    assert b"Choose a PDF file." in unsupported_response.data


@pytest.mark.parametrize(
    ("file_content", "filename", "expected_error"),
    [
        (BytesIO(), "empty.pdf", b"The selected PDF is empty."),
        (BytesIO(b"not a PDF"), "corrupt.pdf", b"not a valid, readable PDF"),
        (pdf_file(()), "no-pages.pdf", b"The selected PDF has no pages."),
    ],
)
def test_pdf_to_word_rejects_invalid_documents(
    client,
    file_content,
    filename,
    expected_error,
):
    response = post_pdf_to_word(client, (file_content, filename))

    assert response.status_code == 400
    assert expected_error in response.data


def test_pdf_to_word_rejects_password_protected_document(client):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=200)
    writer.encrypt("secret")
    protected_pdf = BytesIO()
    writer.write(protected_pdf)
    writer.close()
    protected_pdf.seek(0)

    response = post_pdf_to_word(client, (protected_pdf, "protected.pdf"))

    assert response.status_code == 400
    assert b"Password-protected PDFs cannot be converted." in response.data


def test_large_pdf_to_word_request_returns_clear_error(app):
    app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024
    client = app.test_client()

    response = post_pdf_to_word(
        client,
        (BytesIO(b"x" * (1024 * 1024)), "large.pdf"),
    )

    assert response.status_code == 413
    assert b"The PDF is too large." in response.data
    assert b"maximum upload size is 1 MB" in response.data


@pytest.mark.parametrize(
    "path",
    [
        "/tools/split-pdf",
        "/tools/compress-pdf",
        "/tools/extract-pdf-pages",
        "/tools/rotate-pdf",
        "/tools/images-to-pdf",
        "/tools/word-to-pdf",
    ],
)
def test_new_pdf_tool_pages_are_available(client, path):
    response = client.get(path)

    assert response.status_code == 200
    assert b"<form" in response.data


def test_split_pdf_produces_a_zip_with_requested_ranges(client, app):
    response = client.post(
        "/tools/split-pdf",
        data={
            "pdf": (pdf_file(((100, 200), (300, 400), (500, 600))), "source.pdf"),
            "page_ranges": "1-2,3",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    download = client.get(operation_download_url(response, "split-pdf"))
    with zipfile.ZipFile(BytesIO(download.data)) as archive:
        assert archive.namelist() == ["pages-1-2.pdf", "pages-3-3.pdf"]
        first = PdfReader(BytesIO(archive.read("pages-1-2.pdf")), strict=True)
        second = PdfReader(BytesIO(archive.read("pages-3-3.pdf")), strict=True)
    assert len(first.pages) == 2
    assert len(second.pages) == 1
    other_client = app.test_client()
    assert (
        other_client.get(operation_download_url(response, "split-pdf")).status_code
        == 404
    )


def test_split_pdf_rejects_overlapping_or_out_of_range_pages(client):
    response = client.post(
        "/tools/split-pdf",
        data={
            "pdf": (pdf_file(((100, 200), (300, 400))), "source.pdf"),
            "page_ranges": "1-2,2",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 400
    assert b"Page ranges cannot include a page more than once." in response.data


def test_new_pdf_tool_reports_upload_limit(app):
    app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024
    client = app.test_client()
    response = client.post(
        "/tools/split-pdf",
        data={
            "pdf": (BytesIO(b"x" * (1024 * 1024)), "large.pdf"),
            "page_ranges": "1",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 413
    assert b"maximum total upload size is 1 MB" in response.data


def test_compress_pdf_returns_a_readable_document(client):
    response = client.post(
        "/tools/compress-pdf",
        data={"pdf": (pdf_file(((100, 200), (300, 400))), "source.pdf")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    download = client.get(operation_download_url(response, "compress-pdf"))
    assert download.status_code == 200
    assert len(PdfReader(BytesIO(download.data), strict=True).pages) == 2


def test_extract_pdf_pages_returns_only_selected_pages(client):
    response = client.post(
        "/tools/extract-pdf-pages",
        data={
            "pdf": (pdf_file(((100, 200), (300, 400), (500, 600))), "source.pdf"),
            "page_ranges": "1,3",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    download = client.get(operation_download_url(response, "extract-pdf-pages"))
    extracted = PdfReader(BytesIO(download.data), strict=True)
    assert [float(page.mediabox.width) for page in extracted.pages] == [100, 500]


def test_rotate_pdf_rotates_selected_pages(client):
    response = client.post(
        "/tools/rotate-pdf",
        data={
            "pdf": (pdf_file(((100, 200), (300, 400))), "source.pdf"),
            "page_ranges": "2",
            "rotation": "90",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    download = client.get(operation_download_url(response, "rotate-pdf"))
    rotated = PdfReader(BytesIO(download.data), strict=True)
    assert rotated.pages[0].get("/Rotate", 0) == 0
    assert rotated.pages[1].get("/Rotate", 0) == 90


def test_images_to_pdf_creates_one_page_per_image_in_order(client):
    response = client.post(
        "/tools/images-to-pdf",
        data={
            "images": [
                (image_file((100, 60)), "first.png"),
                (image_file((40, 80), "JPEG"), "second.jpg"),
            ]
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    download = client.get(operation_download_url(response, "images-to-pdf"))
    images_pdf = PdfReader(BytesIO(download.data), strict=True)
    assert len(images_pdf.pages) == 2
    assert float(images_pdf.pages[0].mediabox.width) > float(
        images_pdf.pages[1].mediabox.width
    )


@pytest.mark.skipif(
    sys.platform != "win32"
    or not Path(
        r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE"
    ).is_file(),
    reason="Local Microsoft Word is not installed.",
)
def test_word_to_pdf_uses_local_word_and_downloads_a_valid_pdf(client):
    document = Document()
    document.add_paragraph("Local Word conversion test")
    docx = BytesIO()
    document.save(docx)
    docx.seek(0)

    response = client.post(
        "/tools/word-to-pdf",
        data={"document": (docx, "source.docx")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    download = client.get(operation_download_url(response, "word-to-pdf"))
    converted = PdfReader(BytesIO(download.data), strict=True)
    assert download.status_code == 200
    assert len(converted.pages) >= 1


def test_supplied_test_pdf_merges_and_converts(client):
    pdf_path = Path(__file__).resolve().parents[1] / "test.pdf"
    pdf_data = pdf_path.read_bytes()

    merged_response = post_merge(
        client,
        [
            (BytesIO(pdf_data), "first.pdf"),
            (BytesIO(pdf_data), "second.pdf"),
        ],
    )
    assert merged_response.status_code == 200
    merged_download = client.get(download_url(merged_response))
    assert len(PdfReader(BytesIO(merged_download.data), strict=True).pages) == 2

    converted_response = post_pdf_to_word(client, (BytesIO(pdf_data), "test.pdf"))
    assert converted_response.status_code == 200
    word_download = client.get(pdf_to_word_download_url(converted_response))
    assert len(Document(BytesIO(word_download.data)).sections) == 1

    operations = [
        (
            "split-pdf",
            "/tools/split-pdf",
            {"page_ranges": "1"},
        ),
        ("compress-pdf", "/tools/compress-pdf", {}),
        (
            "extract-pdf-pages",
            "/tools/extract-pdf-pages",
            {"page_ranges": "1"},
        ),
        (
            "rotate-pdf",
            "/tools/rotate-pdf",
            {"page_ranges": "1", "rotation": "90"},
        ),
    ]
    for operation, path, fields in operations:
        response = client.post(
            path,
            data={"pdf": (BytesIO(pdf_data), "test.pdf"), **fields},
            content_type="multipart/form-data",
        )
        assert response.status_code == 200, path
        download = client.get(operation_download_url(response, operation))
        assert download.status_code == 200
        if operation == "split-pdf":
            with zipfile.ZipFile(BytesIO(download.data)) as archive:
                extracted_pdf = PdfReader(
                    BytesIO(archive.read("pages-1-1.pdf")),
                    strict=True,
                )
            assert len(extracted_pdf.pages) == 1
        else:
            output_pdf = PdfReader(BytesIO(download.data), strict=True)
            assert len(output_pdf.pages) == 1
            if operation == "rotate-pdf":
                assert output_pdf.pages[0].get("/Rotate", 0) == 90

    image_path = Path(__file__).resolve().parents[1] / "test.png"
    image_data = image_path.read_bytes()
    images_response = client.post(
        "/tools/images-to-pdf",
        data={"images": [(BytesIO(image_data), "test.png")]},
        content_type="multipart/form-data",
    )
    assert images_response.status_code == 200
    images_download = client.get(
        operation_download_url(images_response, "images-to-pdf")
    )
    assert len(PdfReader(BytesIO(images_download.data), strict=True).pages) == 1
